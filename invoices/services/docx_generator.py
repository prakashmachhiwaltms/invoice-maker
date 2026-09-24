"""Generates a .docx version of the invoice, mirroring the PDF/HTML layout."""
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

BORDER_COLOR = '000000'


def _set_cell_border(cell):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), BORDER_COLOR)
        borders.append(el)
    tc_pr.append(borders)


def _cell_text(cell, text, bold=False, size=9, align=None):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    if align:
        paragraph.alignment = align
    run = paragraph.add_run('' if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(size)
    _set_cell_border(cell)
    return cell


def build_invoice_docx(context: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1)
    section.bottom_margin = Cm(1)

    title = doc.add_heading(context['invoice_title'], level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = header_table.rows[0].cells
    _cell_text(left, f"Company Name: {context['company_name']}", bold=True)
    left.add_paragraph(f"Address: {context['company_address']}")
    left.add_paragraph(f"GSTIN: {context['company_gstin']}")
    _cell_text(right, f"Invoice No. - {context['invoice_number']}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    p = right.add_paragraph(f"Date of Invoice: {context['invoice_date']}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()

    supplier_table = doc.add_table(rows=1, cols=2)
    left, right = supplier_table.rows[0].cells
    _cell_text(left, 'Details of Supplier', bold=True)
    left.add_paragraph(f"Name: {context['supplier_name']}")
    left.add_paragraph(f"Address: {context['supplier_address']}")
    left.add_paragraph(f"State: {context['state']}")
    left.add_paragraph(f"State Code: {context['state_code']}")
    left.add_paragraph(f"Place of Supply: {context['place_of_supply']}")
    left.add_paragraph(f"GSTIN: {context['supplier_gstin']}")

    _cell_text(right, 'Invoice To', bold=True)
    right.add_paragraph(context['invoice_to'] or '-')

    doc.add_paragraph()

    columns = ['Sr No.', 'Description', 'HSN/SAC', 'Qty', 'Taxable Value', 'CGST', 'SGST', 'IGST', 'Amount']
    item_table = doc.add_table(rows=1, cols=len(columns))
    item_table.style = 'Table Grid'
    for i, col in enumerate(columns):
        _cell_text(item_table.rows[0].cells[i], col, bold=True, size=8)

    row = item_table.add_row().cells
    values = [
        '1', context['description'], context['hsn'], str(context['quantity']),
        f"{context['taxable_value']}",
        f"{context['cgst_rate']}% / {context['cgst_amount']}",
        f"{context['sgst_rate']}% / {context['sgst_amount']}",
        f"{context['igst_rate']}% / {context['igst_amount']}",
        f"{context['total']}",
    ]
    for i, val in enumerate(values):
        _cell_text(row[i], val, size=8)

    doc.add_paragraph()

    totals_table = doc.add_table(rows=0, cols=2)
    for label, value in [
        ('Freight', context['freight']),
        ('Insurance', context['insurance']),
        ('Packing', context['packing']),
        ('Total Bill Value (in Figure)', context['total']),
    ]:
        r = totals_table.add_row().cells
        _cell_text(r[0], label, bold=True, size=9)
        _cell_text(r[1], value, size=9, align=WD_ALIGN_PARAGRAPH.RIGHT)

    words_row = totals_table.add_row().cells
    _cell_text(words_row[0], 'Total Bill Value (in Words)', bold=True, size=9)
    _cell_text(words_row[1], context['amount_in_words'], size=9)

    doc.add_paragraph()
    note = doc.add_paragraph(context['reverse_charge_note'])
    note.runs[0].italic = True

    doc.add_paragraph()
    footer = doc.add_paragraph(context['footer_text'])
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].bold = True

    if context.get('signature2') and context['signature2'].image:
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run()
            run.add_picture(context['signature2'].image.path, width=Cm(3))
        except Exception:
            pass

    sig_text = doc.add_paragraph(context['authorized_signatory_text'])
    sig_text.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
